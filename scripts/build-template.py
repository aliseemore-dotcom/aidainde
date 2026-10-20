import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Сборка шаблона: python3 scripts/build-template.py data.json out.xlsx (data.json — из scripts/dump-template-data.ts)
data = json.load(open(sys.argv[1], encoding="utf-8"))
out = sys.argv[2]
INK, TEAL, SOFT, PEACH = "17181A", "1F7A7A", "E3EEEC", "F3DCC8"
wb = Workbook()

# ---------- Инструкция
ws = wb.active; ws.title = "Инструкция"
lines = [
 ("Шаблон каталога «ӘЙДӘ ИНДЕ!»", "title"),
 ("Сюда вносятся места, события и идеи на дом. Из этой таблицы записи загружаются на сайт.", "p"),
 ("", "p"),
 ("Главное правило", "h"),
 ("Вносим только проверенные факты. Ничего не придумываем: нет данных — оставляем ячейку пустой.", "p"),
 ("У каждой записи есть ссылка на первоисточник и дата проверки: когда вы сверили данные с официальной страницей.", "p"),
 ("", "p"),
 ("Как заполнять", "h"),
 ("1. Откройте лист «Места», «События» или «Идеи на дом». Первая строка — заголовки, её не меняйте и не удаляйте.", "p"),
 ("2. Одна строка — одна запись. Колонки со звёздочкой (*) обязательны. Наведите курсор на заголовок — появится подсказка.", "p"),
 ("3. Строку «ПРИМЕР» можно удалить или оставить: она не загрузится.", "p"),
 ("4. Теги выбирайте из листа «Теги», по нескольку через точку с запятой. Нужен минимум один тег-направление.", "p"),
 ("5. Цена: пусто = цена неизвестна (это НЕ «бесплатно»); 0 = бесплатно. Для «Цена указана» выберите из списка.", "p"),
 ("6. Даты: ДД.ММ.ГГГГ, время: ДД.ММ.ГГГГ ЧЧ:ММ — по времени Казани. Пишите как текст, как в примере.", "p"),
 ("7. Если слайг (адрес страницы) оставить пустым, он создастся из названия сам.", "p"),
 ("", "p"),
 ("Что дальше", "h"),
 ("Отправьте файл целиком. Записи будут проверены и загружены как черновики. Публикуются они только после вашей проверки в админке.", "p"),
 ("", "p"),
 ("Чего не делать", "h"),
 ("— не копировать описания с чужих сайтов целиком; пишите своими словами;", "p"),
 ("— не добавлять фото без прав на них (фото вносятся отдельно в админке);", "p"),
 ("— не указывать «бесплатно», если цена просто не найдена;", "p"),
 ("— не вносить события, дату и источник которых вы не проверили.", "p"),
]
for i, (t, kind) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "title": c.font = Font(bold=True, size=16, color=INK)
    elif kind == "h": c.font = Font(bold=True, size=12, color=TEAL)
    else: c.font = Font(size=11, color=INK)
ws.column_dimensions["A"].width = 110

# ---------- Листы данных
def validation(ws, col_idx, formula, title, allow_blank=True, rows=500):
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank, showErrorMessage=True, errorTitle=title, error="Выберите значение из списка")
    ws.add_data_validation(dv); L = get_column_letter(col_idx); dv.add(f"{L}3:{L}{rows}")

# справочник типов мест: 19 названий не помещаются в формулу валидации, поэтому ссылаемся на диапазон листа
kinds_ws = wb.create_sheet("Справочники")
kinds_ws.cell(row=1, column=1, value="Тип места").font = Font(bold=True)
for i, k in enumerate(data["kinds"], 2): kinds_ws.cell(row=i, column=1, value=k["label"])
kinds_ws.column_dimensions["A"].width = 40
KIND_RANGE = f"=Справочники!$A$2:$A${len(data['kinds']) + 1}"
LISTS = {
  "priceUnit": '"с человека,за всех,неизвестно"',
  "setting": '"в помещении,на улице,и там и там,неизвестно"',
  "tri": '"да,нет,неизвестно"',
}
EXAMPLE = {"title": "ПРИМЕР — удалите эту строку или оставьте: она не загрузится"}
thin = Side(style="thin", color="E2E0DB")
for type_, cols in data["columns"].items():
    ws = wb.create_sheet(data["sheets"][type_])
    for i, col in enumerate(cols, 1):
        h = col["header"] + (" *" if col.get("required") else "")
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=INK if col.get("required") else TEAL)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.comment = Comment(col["hint"] or col["header"], "Каталог"); c.comment.width = 320; c.comment.height = 110
        ex = ws.cell(row=2, column=i, value=EXAMPLE.get(col["key"], ""))
        ex.fill = PatternFill("solid", fgColor=PEACH); ex.font = Font(italic=True, color="55575B"); ex.alignment = Alignment(wrap_text=True, vertical="top")
        width = 44 if col["kind"] == "long" else 30 if col["kind"] in ("text", "tags", "hours", "list", "url") else 18
        ws.column_dimensions[get_column_letter(i)].width = width
        if col["kind"] in LISTS: validation(ws, i, LISTS[col["kind"]], col["header"])
        if col["kind"] == "placeKind": validation(ws, i, KIND_RANGE, col["header"])
        if col["kind"] in ("date", "datetime", "tags", "hours", "long", "list", "text", "url"):
            for r in range(3, 501): ws.cell(row=r, column=i).number_format = "@"  # текст: Excel не превращает даты и числа
        if col["kind"] == "number":
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0" if col["key"] not in ("lat", "lng") else "-180", allow_blank=True, showErrorMessage=True, errorTitle=col["header"], error="Нужно число")
            ws.add_data_validation(dv); L = get_column_letter(i); dv.add(f"{L}3:{L}500")
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "B2"
    ws.sheet_view.zoomScale = 100

# ---------- Теги
ws = wb.create_sheet("Теги")
GROUPS = {"category": "Направление и формат", "gameStyle": "Тип игры", "competitionMode": "Азарт", "mood": "Настроение", "energy": "Энергия общения",
          "company": "Для кого", "format": "Формат участия", "physicalLoad": "Нагрузка", "accessibility": "Доступность и ограничения", "setting": "Место"}
for i, h in enumerate(["Тег (пишите именно так)", "Группа", "Заметка"], 1):
    c = ws.cell(row=1, column=i, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=TEAL)
NOTES = {"walks": "Направление: нужен минимум один из шести", "venues": "Направление", "games": "Направление", "culture": "Направление",
         "wellness": "Направление", "home": "Направление (для идей на дом)", "adult_only": "Только для взрослых: такие записи не показываются семьям с детьми",
         "car_required": "Добраться можно только на машине"}
r = 2
for t in data["tags"]:
    ws.cell(row=r, column=1, value=t["label"]); ws.cell(row=r, column=2, value=GROUPS.get(t["group"], t["group"])); ws.cell(row=r, column=3, value=NOTES.get(t["key"], ""))
    r += 1
ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 28; ws.column_dimensions["C"].width = 70; ws.freeze_panes = "A2"
wb.save(out)
print("ok", out)
